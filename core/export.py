"""
Génération du classeur Excel récapitulatif.

Reprise du bloc `export_rapport` de RecrutPro, avec deux changements de fond
(§4 du PROCESS) :

1. L'ALIMENTATION SE FAIT PAR UNE LISTE DE DICTIONNAIRES, plus par des
   requêtes SQL. La fonction devient pure et testable : on lui donne des
   données factices, elle rend un classeur.

2. LE CLASSEUR EST ÉCRIT DANS UN TAMPON MÉMOIRE rendu à l'appelant, jamais
   sur le disque. `openpyxl.Workbook.save()` accepte un objet `BytesIO` aussi
   bien qu'un chemin ; on lui donne le premier. C'est la règle n°1 du §5 du
   cadrage appliquée à la sortie, comme le parser l'applique à l'entrée.

La mise en forme — trois feuilles, palette indigo/ardoise, bandes de couleur
sur les scores — est reprise à l'identique de RecrutPro, comme le prévoit le
§3 du cadrage.

Aucun import de streamlit.
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# Palette — identique à celle de l'export RecrutPro
# ----------------------------------------------------------------------

INDIGO = "4F46E5"
BLANC = "FFFFFF"
GRIS_CLAIR = "F8FAFC"
GRIS_HEADER = "E2E8F0"
GRIS_TEXTE = "64748B"
ARDOISE = "1E293B"

# Bandes de score : vert au-dessus de 70, ambre de 40 à 70, rouge en dessous.
# Ces teintes qualifient la correspondance aux mots-clés, pas le candidat —
# c'est ce que dit la note de bas de tableau, reprise telle quelle.
SCORE_VERT_FOND, SCORE_VERT_TEXTE = "D1FAE5", "065F46"
SCORE_AMBRE_FOND, SCORE_AMBRE_TEXTE = "FEF3C7", "92400E"
SCORE_ROUGE_FOND, SCORE_ROUGE_TEXTE = "FEE2E2", "991B1B"

SEUIL_VERT = 70
SEUIL_AMBRE = 40

SIGNALEMENT_FOND, SIGNALEMENT_TEXTE = "FEF3C7", "92400E"

TYPE_MIME_XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# Texte du §4 du cadrage, garde-fou n°3. Il figure dans l'interface ET dans
# l'export : un classeur circule par courriel, souvent sans la personne qui
# l'a produit, et l'avertissement doit voyager avec les chiffres.
AVERTISSEMENT = (
    "Un score faible signifie que les mots-clés recherchés sont absents du "
    "CV, pas que le candidat est moins qualifié."
)

NOTE_LECTURE = (
    "Le score mesure la correspondance aux mots-clés définis. Il ne mesure "
    "ni la qualité d'un candidat ni son adéquation au poste. Lisez les CV "
    "avant toute décision."
)

NOTE_ILLISIBLES = (
    "Ces CV n'ont pas pu être analysés : ils sont scannés en image, aucun "
    "texte n'en est extractible. Ils ne reçoivent pas de score et ne sont "
    "pas classés. Ouvrez-les directement."
)


# ----------------------------------------------------------------------
# Petits utilitaires de mise en forme
# ----------------------------------------------------------------------


def _styler(cellule, gras=False, couleur=None, fond=None, alignement="left", taille=11):
    """Applique police, fond et alignement à une cellule."""
    cellule.font = Font(bold=gras, color=couleur or ARDOISE, size=taille, name="Calibri")
    if fond:
        cellule.fill = PatternFill("solid", fgColor=fond)
    cellule.alignment = Alignment(
        horizontal=alignement, vertical="center", wrap_text=True
    )


def _border(cellule):
    """Trace un filet fin sur les quatre côtés."""
    filet = Side(style="thin", color=GRIS_HEADER)
    cellule.border = Border(left=filet, right=filet, top=filet, bottom=filet)


def _ecrire_entete(feuille, intitules, hauteur=30):
    """Écrit une ligne d'en-tête indigo et rend son numéro de ligne."""
    feuille.append(intitules)
    ligne = feuille.max_row
    feuille.row_dimensions[ligne].height = hauteur
    for colonne in range(1, len(intitules) + 1):
        cellule = feuille.cell(ligne, colonne)
        _styler(cellule, gras=True, couleur=BLANC, fond=INDIGO, alignement="center")
        _border(cellule)
    return ligne


def _couleurs_du_score(score):
    """Retourne le couple (fond, texte) de la bande de couleur du score."""
    if score >= SEUIL_VERT:
        return SCORE_VERT_FOND, SCORE_VERT_TEXTE
    if score >= SEUIL_AMBRE:
        return SCORE_AMBRE_FOND, SCORE_AMBRE_TEXTE
    return SCORE_ROUGE_FOND, SCORE_ROUGE_TEXTE


def _largeurs(feuille, largeurs):
    for index, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(index)].width = largeur


def _resumer_signalements(signalements):
    """
    Met en forme les signalements d'exclusion pour une cellule de tableau.

    Chaque signalement porte le nom du critère et le ou les termes qui l'ont
    déclenché. Quand les deux sont identiques — cas courant, l'utilisateur
    nommant souvent son critère d'après son mot-clé — on n'écrit pas deux
    fois la même chose.
    """
    if not signalements:
        return "—"
    morceaux = []
    for signalement in signalements:
        critere = signalement.get("critere", "")
        declencheurs = signalement.get("declencheurs") or []
        termes = ", ".join(declencheurs)
        if termes and termes != critere:
            morceaux.append(f"{critere} ({termes})")
        else:
            morceaux.append(critere or termes)
    return " ; ".join(morceaux) or "—"


def _ou_tiret(valeur):
    """Rend la valeur, ou un tiret cadratin si elle est absente."""
    return valeur if valeur else "—"


# ----------------------------------------------------------------------
# Feuille 1 — Résumé
# ----------------------------------------------------------------------


def _feuille_resume(classeur, classes, illisibles, criteres, nom_du_modele, jour):
    feuille = classeur.active
    feuille.title = "Résumé"
    feuille.sheet_view.showGridLines = False
    _largeurs(feuille, [34, 46])

    feuille.row_dimensions[1].height = 50
    feuille.merge_cells("A1:B1")
    _styler(
        feuille["A1"], gras=True, couleur=BLANC, fond=INDIGO,
        alignement="center", taille=16,
    )
    feuille["A1"] = "TriCV — Récapitulatif de tri"

    feuille.merge_cells("A2:B2")
    _styler(feuille["A2"], couleur=GRIS_TEXTE, fond=GRIS_CLAIR, alignement="center", taille=10)
    feuille["A2"] = f"Généré le {jour.strftime('%d/%m/%Y')}"

    feuille.append([])

    nb_signales = sum(1 for candidat in classes if candidat.get("signalements"))
    infos = [
        ("Modèle de poste", nom_du_modele or "Critères personnalisés"),
        ("CV reçus", len(classes) + len(illisibles)),
        ("Analysés et classés", len(classes)),
        ("À examiner manuellement", len(illisibles)),
        ("Signalements", nb_signales),
        ("Critères de sélection", len(criteres)),
    ]

    for intitule, valeur in infos:
        feuille.append([intitule, str(valeur)])
        ligne = feuille.max_row
        feuille.row_dimensions[ligne].height = 25
        _styler(feuille.cell(ligne, 1), gras=True, fond=GRIS_HEADER)
        _styler(feuille.cell(ligne, 2), fond=GRIS_CLAIR)
        _border(feuille.cell(ligne, 1))
        _border(feuille.cell(ligne, 2))

    feuille.append([])

    # L'avertissement méthodologique, sur deux lignes fusionnées pour rester
    # lisible sans que l'utilisateur ait à élargir la colonne.
    for texte, hauteur in ((AVERTISSEMENT, 44), (NOTE_LECTURE, 58)):
        feuille.append([texte])
        ligne = feuille.max_row
        feuille.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=2)
        feuille.row_dimensions[ligne].height = hauteur
        _styler(feuille.cell(ligne, 1), fond=SCORE_AMBRE_FOND, couleur=SCORE_AMBRE_TEXTE)
        _border(feuille.cell(ligne, 1))

    feuille.append([])

    feuille.append(["Traitement éphémère — aucune conservation"])
    ligne = feuille.max_row
    feuille.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=2)
    _styler(feuille.cell(ligne, 1), gras=True, couleur=GRIS_TEXTE)

    feuille.append([
        "Les CV analysés pour produire ce récapitulatif n'ont jamais été "
        "enregistrés sur un disque ni transmis à un tiers. Ce fichier est la "
        "seule trace qui subsiste."
    ])
    ligne = feuille.max_row
    feuille.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=2)
    feuille.row_dimensions[ligne].height = 58
    _styler(feuille.cell(ligne, 1), couleur=GRIS_TEXTE, taille=10)


# ----------------------------------------------------------------------
# Feuille 2 — Classement
# ----------------------------------------------------------------------


def _feuille_classement(classeur, classes, illisibles, criteres):
    feuille = classeur.create_sheet("Classement")
    feuille.sheet_view.showGridLines = False

    noms_criteres = [critere.get("nom", "") for critere in criteres]

    intitules = ["Rang", "Candidat", "Email", "Téléphone", "Score /100", "Signalements"]
    intitules += [f"{nom} (pts)" for nom in noms_criteres]
    intitules += ["Mots-clés trouvés", "Mots-clés absents", "Fichier"]

    _largeurs(
        feuille,
        [6, 28, 30, 20, 12, 22] + [16] * len(noms_criteres) + [46, 46, 30],
    )
    _ecrire_entete(feuille, intitules)

    for rang, candidat in enumerate(classes, start=1):
        score = candidat.get("score", 0)
        detail = {ligne["critere"]: ligne for ligne in candidat.get("detail", [])}

        trouves = []
        absents = []
        for nom in noms_criteres:
            ligne_detail = detail.get(nom)
            if ligne_detail:
                trouves.extend(ligne_detail.get("trouves", []))
                absents.extend(ligne_detail.get("absents", []))

        valeurs = [
            rang,
            _ou_tiret(candidat.get("nom")),
            _ou_tiret(candidat.get("email")),
            _ou_tiret(candidat.get("telephone")),
            score,
            _resumer_signalements(candidat.get("signalements")),
        ]
        for nom in noms_criteres:
            ligne_detail = detail.get(nom)
            if ligne_detail:
                valeurs.append(f"{ligne_detail['points']} / {ligne_detail['poids']}")
            else:
                valeurs.append("—")
        valeurs += [", ".join(trouves) or "—", ", ".join(absents) or "—",
                    candidat.get("fichier", "—")]

        feuille.append(valeurs)
        ligne = feuille.max_row
        feuille.row_dimensions[ligne].height = 25

        fond_ligne = BLANC if rang % 2 == 0 else GRIS_CLAIR
        score_fond, score_texte = _couleurs_du_score(score)

        for colonne in range(1, len(valeurs) + 1):
            cellule = feuille.cell(ligne, colonne)
            _border(cellule)
            if colonne == 1:
                _styler(cellule, fond=fond_ligne, alignement="center", couleur=GRIS_TEXTE)
            elif colonne == 5:
                _styler(cellule, gras=True, couleur=score_texte,
                        fond=score_fond, alignement="center")
            elif colonne == 6 and candidat.get("signalements"):
                _styler(cellule, couleur=SIGNALEMENT_TEXTE,
                        fond=SIGNALEMENT_FOND, alignement="center")
            elif 7 <= colonne <= 6 + len(noms_criteres):
                _styler(cellule, fond=fond_ligne, alignement="center")
            else:
                _styler(cellule, fond=fond_ligne)

    # ------------------------------------------------------------------
    # Bloc « À examiner manuellement »
    # ------------------------------------------------------------------
    #
    # Ces CV sont placés APRÈS le classement et dans un bloc séparé, jamais
    # en fin de tableau avec un score de 0. Les mettre dans le classement
    # affirmerait qu'ils correspondent mal au poste, alors qu'on n'en sait
    # rien : on n'a pas pu les lire (§7.5 du cadrage).
    if not illisibles:
        return

    feuille.append([])
    feuille.append([])

    feuille.append(["À examiner manuellement"])
    ligne = feuille.max_row
    _styler(feuille.cell(ligne, 1), gras=True, taille=13)

    feuille.append([NOTE_ILLISIBLES])
    ligne = feuille.max_row
    feuille.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
    feuille.row_dimensions[ligne].height = 40
    _styler(feuille.cell(ligne, 1), couleur=GRIS_TEXTE, taille=10)

    _ecrire_entete(feuille, ["Fichier", "Taille", "Raison", "Score /100"])

    for index, candidat in enumerate(illisibles, start=1):
        feuille.append([
            candidat.get("fichier", "—"),
            _ou_tiret(candidat.get("taille")),
            "Document scanné, aucun texte lisible",
            "—",
        ])
        ligne = feuille.max_row
        feuille.row_dimensions[ligne].height = 25
        fond = BLANC if index % 2 == 0 else GRIS_CLAIR
        for colonne in range(1, 5):
            cellule = feuille.cell(ligne, colonne)
            _border(cellule)
            _styler(cellule, fond=fond, alignement="center" if colonne == 4 else "left")


# ----------------------------------------------------------------------
# Feuille 3 — Critères
# ----------------------------------------------------------------------


def _feuille_criteres(classeur, criteres):
    feuille = classeur.create_sheet("Critères")
    feuille.sheet_view.showGridLines = False
    _largeurs(feuille, [28, 14, 14, 62])

    _ecrire_entete(feuille, ["Critère", "Type", "Poids", "Mots-clés recherchés"])

    for index, critere in enumerate(criteres, start=1):
        groupes = critere.get("groupes", [])
        # On réaffiche les mots-clés tels que l'utilisateur les a saisis
        # quand on les a sous la main, sinon on les reconstruit depuis les
        # groupes. Le lecteur du classeur doit pouvoir rejouer le tri.
        mots_cles = critere.get("mots_cles")
        if not mots_cles:
            mots_cles = "\n".join(
                groupe[0] if len(groupe) == 1 else f"{groupe[0]} = {', '.join(groupe[1:])}"
                for groupe in groupes
            )

        type_critere = critere.get("type", "requis")
        feuille.append([
            critere.get("nom", ""),
            "Exclusion" if type_critere == "exclusion" else "Requis",
            "—" if type_critere == "exclusion" else critere.get("poids", 0),
            mots_cles or "—",
        ])
        ligne = feuille.max_row
        # Une ligne par groupe de mots-clés, pour que rien ne soit tronqué.
        feuille.row_dimensions[ligne].height = max(25, 15 * max(1, len(groupes)))
        fond = BLANC if index % 2 == 0 else GRIS_CLAIR
        for colonne in range(1, 5):
            cellule = feuille.cell(ligne, colonne)
            _border(cellule)
            _styler(cellule, fond=fond, alignement="center" if colonne in (2, 3) else "left")

    feuille.append([])
    feuille.append([
        "Un groupe de synonymes est écrit « terme principal = synonyme, synonyme ». "
        "Le groupe compte pour trouvé si au moins un de ses termes figure dans le CV."
    ])
    ligne = feuille.max_row
    feuille.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
    feuille.row_dimensions[ligne].height = 30
    _styler(feuille.cell(ligne, 1), couleur=GRIS_TEXTE, taille=10)


# ----------------------------------------------------------------------
# Point d'entrée
# ----------------------------------------------------------------------


def generer_classeur(candidats, criteres, nom_du_modele=None, jour=None):
    """
    Construit le classeur récapitulatif et le rend dans un tampon mémoire.

    `candidats` : la liste des résultats, telle que produite par le parser
                  puis le scorer. Chaque entrée porte au moins `fichier`,
                  `lisible`, `nom`, `email`, `telephone`, et, si elle est
                  lisible, `score`, `signalements` et `detail`.
    `criteres`  : les critères utilisés, avec `nom`, `poids`, `type` et
                  `groupes` (et `mots_cles` si on l'a).
    `nom_du_modele` : l'intitulé du modèle choisi, ou None si les critères
                  ont été saisis à la main.
    `jour`      : la date à inscrire. Laissée à None, c'est aujourd'hui.

    POURQUOI `jour` EST UN PARAMÈTRE. Lire l'horloge rend une fonction
    impure : deux appels identiques peuvent produire deux fichiers
    différents, et un test ne peut rien affirmer sur le contenu. En laissant
    l'appelant fournir la date, la fonction redevient vérifiable, tout en
    restant commode à l'usage puisque None suffit.

    Retourne un `io.BytesIO` positionné au début, prêt pour
    `st.download_button`. Aucun fichier n'est écrit sur le disque.
    """
    if jour is None:
        jour = date.today()

    # Les critères d'exclusion ne produisent pas de colonne de points : ils
    # ne rapportent rien (§7.4). Ils restent listés dans la feuille 3.
    criteres_requis = [
        critere for critere in criteres
        if critere.get("type", "requis") != "exclusion"
    ]

    classes = sorted(
        (candidat for candidat in candidats if candidat.get("lisible")),
        key=lambda candidat: candidat.get("score", 0),
        reverse=True,
    )
    illisibles = [candidat for candidat in candidats if not candidat.get("lisible")]

    classeur = Workbook()
    _feuille_resume(classeur, classes, illisibles, criteres, nom_du_modele, jour)
    _feuille_classement(classeur, classes, illisibles, criteres_requis)
    _feuille_criteres(classeur, criteres)

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


def nom_du_fichier_export(jour=None):
    """
    Nom proposé au téléchargement.

    Volontairement NEUTRE : ni nom de candidat, ni intitulé de poste. Un nom
    de fichier se retrouve dans l'historique de téléchargement du navigateur
    et dans les pièces jointes des courriels — c'est-à-dire à des endroits
    durables que l'utilisateur ne contrôle pas.
    """
    if jour is None:
        jour = date.today()
    return f"tricv_recapitulatif_{jour.strftime('%Y-%m-%d')}.xlsx"
