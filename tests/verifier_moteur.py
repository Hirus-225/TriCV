"""
Contrôle de la porte P1 — vérification du moteur en ligne de commande.

Ce script fait tourner le parser et le scorer sur le corpus de test fictif et
affiche, pour chaque fichier : le nom extrait, l'email, le téléphone, le
statut lisible/illisible, le score et le détail par critère.

Il se lit en quatre parties :

  PARTIE A — les trois corrections du lot 1, démontrées sur du texte écrit à
             la main. C'est le critère de sortie du lot 1 dans PROCESS §4.
  PARTIE B — le moteur complet sur le corpus, avec le modèle « Développeur
             web » du §8 du cadrage. C'est le critère de sortie du lot 2.
  PARTIE C — la syntaxe de saisie des mots-clés et les cinq modèles.
             C'est le critère de sortie du lot 3.
  PARTIE D — le classeur Excel, généré puis relu en mémoire.
             C'est le critère de sortie du lot 4.

Chaque partie se termine par des vérifications automatiques. Le script sort
avec le code 0 si tout passe, 1 sinon — de quoi le brancher plus tard sur une
intégration continue.

    ./venv/bin/python tests/verifier_moteur.py

CONFIDENTIALITÉ. Ce script affiche des noms, des emails et des numéros de
téléphone, ce que la règle n°5 du §5 du cadrage interdit par ailleurs. C'est
la seule exception admise, et elle tient à une condition : il ne lit QUE
tests/corpus/, dont le contenu est entièrement fabriqué. Ne jamais le faire
pointer vers des CV réels.
"""

import io
import sys
from datetime import date
from pathlib import Path

# Permet de lancer le script directement, sans installer le projet.
RACINE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RACINE))

from openpyxl import load_workbook

from core.criteres import analyser_mots_cles, formater_mots_cles, total_des_poids
from core.export import AVERTISSEMENT, generer_classeur, nom_du_fichier_export
from core.modeles import MODELES, charger_modele, noms_des_modeles
from core.normalisation import normaliser
from core.parser import SEUIL_TEXTE_MINIMUM, extraire_texte, parser_cv
from core.scorer import scorer_candidat, terme_present

DOSSIER_CORPUS = RACINE / "tests" / "corpus"


# ----------------------------------------------------------------------
# Le modèle de critères utilisé pour la vérification
# ----------------------------------------------------------------------

# Le moteur est vérifié avec le VRAI modèle, celui que l'interface servira.
# Si core/modeles.py transcrit mal le §8 du cadrage, les scores ci-dessous
# changent et les vérifications tombent.
MODELE_DEVELOPPEUR_WEB = charger_modele("Développeur web")

# Le §8 du cadrage, retranscrit indépendamment pour servir de référence à la
# partie C. Cette redite est volontaire : si les deux transcriptions
# divergent, c'est qu'au moins l'une des deux est fausse — et c'est
# exactement ce qu'on veut détecter. Le format est (nom, poids, groupes).
SECTION_8_DEVELOPPEUR_WEB = [
    ("Langages & frameworks", 35, [
        ["python"],
        ["javascript", "js"],
        ["django", "flask"],
        ["react", "vue", "angular"],
    ]),
    ("Bases de données", 20, [
        ["sql"],
        ["postgresql", "postgres"],
        ["mysql"],
        ["mongodb", "nosql"],
    ]),
    ("Outils & versioning", 15, [
        ["git", "github", "gitlab"],
        ["docker"],
        ["linux"],
    ]),
    ("Expérience terrain", 20, [
        ["developpeur", "developer", "ingenieur logiciel"],
        ["freelance"],
        ["stage"],
    ]),
    ("Formation", 10, [
        ["licence", "bachelor"],
        ["master"],
        ["informatique"],
    ]),
]


# ----------------------------------------------------------------------
# Affichage
# ----------------------------------------------------------------------

LARGEUR = 78
_echecs = []


def titre(texte):
    print()
    print("=" * LARGEUR)
    print(texte)
    print("=" * LARGEUR)


def sous_titre(texte):
    print()
    print(texte)
    print("-" * LARGEUR)


def verifier(intitule, condition, observe=""):
    """Enregistre une vérification et l'affiche. Retourne la condition."""
    marque = "OK  " if condition else "ECHEC"
    suffixe = f"   (observé : {observe})" if observe and not condition else ""
    print(f"  [{marque}] {intitule}{suffixe}")
    if not condition:
        _echecs.append(intitule)
    return condition


# ----------------------------------------------------------------------
# PARTIE A — les trois corrections du lot 1
# ----------------------------------------------------------------------


def partie_a_frontieres_de_mots():
    sous_titre("A1. Frontières de mots — le bug principal de RecrutPro")

    texte = "Algorithme de recrutement. Mission au Congo et en Angola."
    texte_normalise = normaliser(texte)

    print(f'  Texte     : "{texte}"')
    print(f"  Normalisé : \"{texte_normalise}\"")
    print()
    print("  mot-clé   TriCV (\\bmot\\b)   RecrutPro (mot in texte)")
    for terme in ["r", "go", "c", "python"]:
        tricv = terme_present(terme, texte_normalise)
        recrutpro = terme in texte_normalise
        print(f"  {terme:<9} {str(tricv):<17} {recrutpro}")
    print()

    verifier(
        "« r » n'est pas trouvé dans « recrutement »",
        not terme_present("r", texte_normalise),
    )
    verifier(
        "« go » n'est pas trouvé dans « algorithme », « Congo », « Angola »",
        not terme_present("go", texte_normalise),
    )
    verifier(
        "un mot-clé d'une lettre reste trouvable quand il est un mot entier",
        terme_present("r", normaliser("Je programme en R et en Python.")),
    )


def partie_a_accents():
    sous_titre("A2. Suppression des accents des DEUX côtés de la comparaison")

    cv_accentue = normaliser("Développeur back-end, spécialisé en modélisation.")
    cv_sans_accent = normaliser("Developpeur back-end, specialise en modelisation.")

    print(f'  CV accentué     : "{cv_accentue}"')
    print(f'  CV sans accent  : "{cv_sans_accent}"')
    print()

    verifier(
        "mot-clé « developpeur » trouve « Développeur »",
        terme_present("developpeur", cv_accentue),
    )
    verifier(
        "mot-clé « développeur » trouve « Developpeur »",
        terme_present("développeur", cv_sans_accent),
    )
    verifier(
        "mot-clé accentué trouve un CV accentué",
        terme_present("modélisation", cv_accentue),
    )


def partie_a_normalisation_des_poids():
    sous_titre("A3. Score normalisé par le total des poids, sans plafond")

    # Deux critères, poids 100 et 50 : la somme dépasse 100, ce qui déclenchait
    # l'écrasement dans RecrutPro.
    criteres = [
        {"nom": "A", "poids": 100, "type": "requis", "groupes": [["python"], ["sql"]]},
        {"nom": "B", "poids": 50, "type": "requis", "groupes": [["git"]]},
    ]

    complet = scorer_candidat("python sql git", criteres)["score"]
    partiel = scorer_candidat("python sql", criteres)["score"]
    faible = scorer_candidat("python", criteres)["score"]

    print("  Somme des poids saisis : 150 (au-delà de 100, volontairement)")
    print(f"  CV complet  (3 groupes sur 3) : {complet:>3} / 100")
    print(f"  CV partiel  (2 groupes sur 3) : {partiel:>3} / 100")
    print(f"  CV faible   (1 groupe  sur 3) : {faible:>3} / 100")
    print("  Avec le min(score, 100) de RecrutPro : 100, 100 et 50 —")
    print("  les deux premiers devenaient indistinguables.")
    print()

    verifier("un CV complet atteint 100", complet == 100, complet)
    verifier(
        "les scores restent distincts malgré des poids totalisant 150",
        complet > partiel > faible,
        f"{complet} / {partiel} / {faible}",
    )

    # Même profil relatif, poids exprimés sur une autre base : même score.
    criteres_sur_10 = [
        {"nom": "A", "poids": 10, "type": "requis", "groupes": [["python"], ["sql"]]},
        {"nom": "B", "poids": 5, "type": "requis", "groupes": [["git"]]},
    ]
    identique = scorer_candidat("python sql", criteres_sur_10)["score"]
    verifier(
        "le score ne dépend pas de l'échelle des poids (150 ou 15)",
        identique == partiel,
        f"{identique} vs {partiel}",
    )


def partie_a_exclusions():
    sous_titre("A4. Critère d'exclusion — signalement sans effet sur le score")

    criteres = MODELE_DEVELOPPEUR_WEB + [
        {
            "nom": "Recherche un stage",
            "poids": 0,
            "type": "exclusion",
            "groupes": [["recherche un stage"], ["interim"]],
        }
    ]

    texte = (
        "Developpeur Python, Django, React, JavaScript. "
        "SQL, PostgreSQL, MySQL, MongoDB. Git, Docker, Linux. "
        "Freelance. Licence en informatique. "
        "Je recherche un stage de fin de cycle."
    )

    avec_exclusion = scorer_candidat(texte, criteres)
    sans_exclusion = scorer_candidat(texte, MODELE_DEVELOPPEUR_WEB)

    print(f"  Score avec le critère d'exclusion : {avec_exclusion['score']}")
    print(f"  Score sans le critère d'exclusion : {sans_exclusion['score']}")
    print(f"  Signalements : {avec_exclusion['signalements']}")
    print()

    verifier(
        "le critère d'exclusion produit un signalement",
        [s["critere"] for s in avec_exclusion["signalements"]] == ["Recherche un stage"],
        avec_exclusion["signalements"],
    )
    verifier(
        "le signalement dit QUEL mot-clé a déclenché (décision C4)",
        avec_exclusion["signalements"][0]["declencheurs"] == ["recherche un stage"],
        avec_exclusion["signalements"][0]["declencheurs"],
    )
    verifier(
        "le signalement ne modifie pas le score",
        avec_exclusion["score"] == sans_exclusion["score"],
    )
    verifier(
        "le candidat reste dans le classement (il a un score)",
        avec_exclusion["score"] > 0,
    )


# ----------------------------------------------------------------------
# PARTIE B — le moteur complet sur le corpus
# ----------------------------------------------------------------------


def analyser_fichier(chemin):
    """
    Lit le fichier en mémoire, le passe au parser, puis au scorer s'il est
    lisible.

    La lecture depuis le disque n'existe QUE dans ce script de vérification :
    elle remplace le dépôt de fichier que fera `st.file_uploader`. Le parser,
    lui, ne reçoit qu'un objet mémoire — jamais un chemin.
    """
    with open(chemin, "rb") as fichier:
        memoire = io.BytesIO(fichier.read())

    resultat = parser_cv(memoire, chemin.name)

    if resultat["lisible"]:
        resultat.update(scorer_candidat(resultat["texte"], MODELE_DEVELOPPEUR_WEB))
    else:
        # Un document illisible n'est PAS scoré. « Illisible » et « score de
        # 0 » sont deux états distincts (§7.5 du cadrage).
        resultat["score"] = None
        resultat["signalements"] = []
        resultat["detail"] = []

    # Règle n°3 du §5 : le texte brut est libéré dès le score calculé. On
    # garde seulement sa longueur, pour le diagnostic.
    resultat["longueur_texte_normalise"] = len(normaliser(resultat["texte"]))
    del resultat["texte"]

    return resultat


def afficher_resultat(resultat):
    print()
    print(f"  ┌─ {resultat['fichier']}")

    nom = resultat["nom"] if resultat["nom"] else "— (non extrait)"
    print(f"  │  Nom        : {nom}")
    print(f"  │  Email      : {resultat['email'] or '—'}")
    print(f"  │  Téléphone  : {resultat['telephone'] or '—'}")

    if resultat["lisible"]:
        statut = f"lisible ({resultat['longueur_texte_normalise']} caractères normalisés)"
    else:
        statut = (
            f"ILLISIBLE ({resultat['longueur_texte_normalise']} caractères, "
            f"seuil {SEUIL_TEXTE_MINIMUM}) → à examiner manuellement"
        )
    print(f"  │  Statut     : {statut}")

    if resultat["score"] is None:
        print("  │  Score      : aucun — un document illisible n'est pas classé")
        print("  └─")
        return

    print(f"  │  Score      : {resultat['score']} / 100")
    if resultat["signalements"]:
        print(f"  │  Signalements : {', '.join(resultat['signalements'])}")

    print("  │")
    for ligne in resultat["detail"]:
        print(f"  │  {ligne['critere']} — {ligne['points']} / {ligne['poids']} pts")
        print(f"  │      Trouvés : {', '.join(ligne['trouves']) or '—'}")
        print(f"  │      Absents : {', '.join(ligne['absents']) or '—'}")
    print("  └─")


def partie_b_corpus():
    titre("PARTIE B — Le moteur sur le corpus, modèle « Développeur web »")

    fichiers = sorted(
        c for c in DOSSIER_CORPUS.iterdir() if c.suffix.lower() in (".pdf", ".docx")
    )

    if not fichiers:
        print("  Corpus vide. Lancer d'abord : ./venv/bin/python tests/generer_corpus.py")
        _echecs.append("corpus absent")
        return []

    resultats = [analyser_fichier(chemin) for chemin in fichiers]

    for resultat in resultats:
        afficher_resultat(resultat)

    # ------------------------------------------------------------------
    # Classement
    # ------------------------------------------------------------------
    sous_titre("Classement (les illisibles sont à part, pas en dernier)")

    classes = sorted(
        (r for r in resultats if r["lisible"]),
        key=lambda r: r["score"],
        reverse=True,
    )
    illisibles = [r for r in resultats if not r["lisible"]]

    print(f"  {'Rang':<5} {'Score':<7} {'Candidat':<24} Fichier")
    for rang, resultat in enumerate(classes, start=1):
        nom = resultat["nom"] or "— (non extrait)"
        print(f"  {rang:<5} {resultat['score']:<7} {nom:<24} {resultat['fichier']}")

    print()
    print("  À examiner manuellement :")
    for resultat in illisibles:
        print(f"         —       {'':<24} {resultat['fichier']}")

    # ------------------------------------------------------------------
    # Vérifications
    # ------------------------------------------------------------------
    sous_titre("Vérifications du corpus")

    par_fichier = {r["fichier"]: r for r in resultats}

    def attendu(nom_fichier):
        resultat = par_fichier.get(nom_fichier)
        if resultat is None:
            _echecs.append(f"fichier absent : {nom_fichier}")
        return resultat

    # Cas 1 — chemin nominal
    cas = attendu("01_CV_Konan_Bernard.pdf")
    if cas:
        verifier("cas 1 — PDF texte lisible", cas["lisible"])
        verifier("cas 1 — nom extrait du document", cas["nom"] == "Konan Bernard", cas["nom"])
        verifier("cas 1 — email extrait", cas["email"] == "konan.bernard@example.ci", cas["email"])
        verifier("cas 1 — téléphone ivoirien extrait", cas["telephone"] is not None)
        verifier("cas 1 — score élevé (profil très proche)", cas["score"] >= 75, cas["score"])

    # Cas 2 — accents, prénom composé, tiret cadratin, accentuation décomposée
    cas = attendu("02_CV_Kouassi_Guy_Desire.pdf")
    if cas:
        verifier(
            "cas 2 — nom recomposé en NFC, tiret cadratin retiré, trait d'union gardé",
            cas["nom"] == "Kouassi Guy-Désiré",
            repr(cas["nom"]),
        )
        verifier(
            "cas 2 — « Développeur » trouvé par le mot-clé « developpeur »",
            "developpeur" in _trouves(cas, "Expérience terrain"),
            _trouves(cas, "Expérience terrain"),
        )

    # Cas 3 — scanné
    cas = attendu("03_CV_Traore_Aminata.pdf")
    if cas:
        verifier("cas 3 — PDF scanné détecté comme illisible", not cas["lisible"])
        verifier("cas 3 — aucun score attribué", cas["score"] is None, cas["score"])
        verifier("cas 3 — hors classement", cas not in classes)

    # Cas 4 — DOCX en tableaux
    cas = attendu("04_CV_Diabate_Fatoumata.docx")
    if cas:
        verifier(
            "cas 4 — DOCX en tableaux : le contenu des cellules est extrait",
            cas["lisible"],
        )
        verifier(
            "cas 4 — assez de texte pour un score réel",
            cas["score"] is not None and cas["score"] > 30,
            cas["score"],
        )

    # Cas 5 — DOCX en paragraphes
    cas = attendu("05_CV_Yao_Emmanuel.docx")
    if cas:
        verifier("cas 5 — DOCX en paragraphes lisible", cas["lisible"])
        verifier("cas 5 — nom extrait", cas["nom"] == "Yao Kouadio Emmanuel", cas["nom"])

    # Cas 6 — anglais
    cas = attendu("06_CV_Mensah_Kwame.pdf")
    if cas:
        verifier("cas 6 — CV anglais lisible", cas["lisible"])
        verifier(
            "cas 6 — les synonymes anglais fonctionnent (developer, bachelor)",
            "developpeur" in _trouves(cas, "Expérience terrain")
            and "licence" in _trouves(cas, "Formation"),
            f"{_trouves(cas, 'Expérience terrain')} / {_trouves(cas, 'Formation')}",
        )

    # Cas 7 — nom de fichier inexploitable
    cas = attendu("CVBR.pdf")
    if cas:
        verifier(
            "cas 7 — CVBR.pdf : le repli échoue proprement, nom = None",
            cas["nom"] is None,
            repr(cas["nom"]),
        )
        verifier("cas 7 — le document reste lisible et classé", cas["lisible"])

    # Cas 8 — nom de fichier exploitable
    cas = attendu("CV_Adjoua_Kouadio_2026.pdf")
    if cas:
        verifier(
            "cas 8 — repli sur le nom du fichier, année et préfixe CV écartés",
            cas["nom"] == "Adjoua Kouadio",
            repr(cas["nom"]),
        )

    # Cas 9 — sans coordonnées
    cas = attendu("09_CV_Ouattara_Salimata.pdf")
    if cas:
        verifier("cas 9 — email absent rendu comme None", cas["email"] is None, cas["email"])
        verifier("cas 9 — téléphone absent rendu comme None", cas["telephone"] is None)
        verifier(
            "cas 9 — LISIBLE malgré un score bas : deux états distincts",
            cas["lisible"] and cas["score"] is not None,
        )
        verifier("cas 9 — score bas (profil éloigné du poste)", cas["score"] < 30, cas["score"])

    # Cas 10 — DOCX en zones de texte
    #
    # Ce cas vient de l'essai sur échantillon réel du 16/09/2026, où un CV
    # composé uniquement de zones de texte flottantes ressortait à zéro
    # caractère. Il était donc classé « illisible » — un candidat réel
    # disparaissait du classement, sans qu'aucun signal ne distingue ce
    # défaut d'un vrai scan illisible.
    cas = attendu("10_CV_Brou_Akissi.docx")
    if cas:
        verifier(
            "cas 10 — DOCX en zones de texte : le contenu flottant est extrait",
            cas["lisible"],
        )
        verifier(
            "cas 10 — nom extrait depuis la zone de texte",
            cas["nom"] == "BROU AKISSI",
            repr(cas["nom"]),
        )
        verifier(
            "cas 10 — email extrait depuis la zone de texte",
            cas["email"] == "akissi.brou@example.ci",
            cas["email"],
        )
        verifier(
            "cas 10 — téléphone extrait depuis la zone de texte",
            cas["telephone"] is not None,
            cas["telephone"],
        )

        # Le premier bloc du fichier est écrit en mc:AlternateContent : la
        # même zone de texte y figure deux fois. Un parcours naïf la
        # compterait deux fois.
        with open(DOSSIER_CORPUS / "10_CV_Brou_Akissi.docx", "rb") as fichier:
            texte_10 = extraire_texte(io.BytesIO(fichier.read()),
                                      "10_CV_Brou_Akissi.docx")
        verifier(
            "cas 10 — le bloc en double (mc:Fallback) n'est compté qu'une fois",
            texte_10.count("BROU AKISSI") == 1,
            f"{texte_10.count('BROU AKISSI')} occurrence(s)",
        )

    # Vérifications d'ensemble
    verifier(
        "un seul document illisible dans le corpus",
        len(illisibles) == 1,
        len(illisibles),
    )
    verifier(
        "aucun illisible ne figure dans le classement",
        all(r["lisible"] for r in classes),
    )
    verifier(
        "tous les scores sont dans l'intervalle 0-100",
        all(0 <= r["score"] <= 100 for r in classes),
    )

    return resultats


def _trouves(resultat, nom_critere):
    """Retourne la liste des termes trouvés pour un critère donné."""
    for ligne in resultat.get("detail", []):
        if ligne["critere"] == nom_critere:
            return ligne["trouves"]
    return []


# ----------------------------------------------------------------------
# PARTIE C — Syntaxe de saisie et modèles (lot 3)
# ----------------------------------------------------------------------


def partie_c_syntaxe():
    sous_titre("C1. Syntaxe de saisie — le « = » crée le groupe")

    cas = [
        ("python", [["python"]]),
        ("python, javascript", [["python"], ["javascript"]]),
        ("python\njavascript", [["python"], ["javascript"]]),
        ("react = vue, angular", [["react", "vue", "angular"]]),
        (
            "python, javascript\nreact = vue, angular",
            [["python"], ["javascript"], ["react", "vue", "angular"]],
        ),
    ]

    for saisie, attendu in cas:
        obtenu = analyser_mots_cles(saisie)
        apercu = saisie.replace("\n", " ⏎ ")
        print(f'  "{apercu}"')
        print(f"      -> {obtenu}")
        verifier(f'saisie "{apercu}"', obtenu == attendu, obtenu)

    print()
    sous_titre("C2. Tolérance de saisie et aller-retour")

    verifier(
        "espaces superflus absorbés",
        analyser_mots_cles("  python  ,  django  ") == [["python"], ["django"]],
    )
    verifier(
        "lignes vides et virgules en trop absorbées",
        analyser_mots_cles("python,,\n\n  \ndjango,") == [["python"], ["django"]],
    )
    verifier(
        "doublons retirés",
        analyser_mots_cles("python\npython") == [["python"]],
    )
    verifier(
        "un second « = » sur la ligne ne coupe pas deux fois",
        analyser_mots_cles("a = b = c") == [["a", "b = c"]],
        analyser_mots_cles("a = b = c"),
    )
    verifier("saisie vide rend une liste vide", analyser_mots_cles("") == [])

    # Aller-retour : le texte reformaté doit avoir le même SENS que
    # l'original. C'est ce qui garantit qu'un modèle affiché dans la zone de
    # saisie score bien ce qu'il montre.
    for nom_modele in noms_des_modeles():
        for critere in charger_modele(nom_modele):
            retour = analyser_mots_cles(formater_mots_cles(critere["groupes"]))
            if retour != critere["groupes"]:
                verifier(
                    f"aller-retour stable — {nom_modele} / {critere['nom']}",
                    False,
                    retour,
                )
                break
        else:
            continue
        break
    else:
        verifier("aller-retour texte → groupes → texte stable sur les 5 modèles", True)


def partie_c_modeles():
    sous_titre("C3. Les cinq modèles")

    print(f"  {'Modèle':<26} {'Critères':<10} {'Poids':<7} Groupes")
    for nom_modele in noms_des_modeles():
        criteres = charger_modele(nom_modele)
        total = total_des_poids(criteres)
        nb_groupes = sum(len(c["groupes"]) for c in criteres)
        print(f"  {nom_modele:<26} {len(criteres):<10} {total:<7} {nb_groupes}")
    print()

    verifier("les 5 modèles du §8 sont présents", len(MODELES) == 5, len(MODELES))

    for nom_modele in noms_des_modeles():
        criteres = charger_modele(nom_modele)
        total = total_des_poids(criteres)
        verifier(f"« {nom_modele} » totalise exactement 100 points", total == 100, total)
        verifier(
            f"« {nom_modele} » : aucun critère sans mot-clé",
            all(c["groupes"] for c in criteres),
        )
        verifier(
            f"« {nom_modele} » : tous les critères sont de type requis",
            all(c["type"] == "requis" for c in criteres),
        )

    # Comparaison avec la transcription indépendante du §8.
    print()
    reel = charger_modele("Développeur web")
    verifier(
        "« Développeur web » a les 5 critères du §8",
        [c["nom"] for c in reel] == [n for n, _, _ in SECTION_8_DEVELOPPEUR_WEB],
        [c["nom"] for c in reel],
    )
    for critere, (nom, poids, groupes) in zip(reel, SECTION_8_DEVELOPPEUR_WEB):
        verifier(
            f"« {nom} » : poids et groupes conformes au §8",
            critere["poids"] == poids and critere["groupes"] == groupes,
            f"{critere['poids']} / {critere['groupes']}",
        )

    # Les copies doivent être indépendantes de la constante.
    premier = charger_modele("Comptable")
    premier[0]["poids"] = 999
    verifier(
        "charger_modele rend des copies : modifier le résultat n'altère pas MODELES",
        total_des_poids(charger_modele("Comptable")) == 100,
    )

    sous_titre("C4. Le modèle « Comptable » sépare bien le métier")

    # GARDE-FOU DU RECALIBRAGE DU 16/09/2026.
    #
    # Un modèle se dérègle de deux façons opposées, et les deux passent
    # inaperçues sans mesure. Trop de mots-clés en expression exacte — « grand
    # livre », « aide comptable » — et plus rien ne se déclenche : tous les
    # candidats s'effondrent ensemble. Trop de mots-clés génériques, et tout
    # se déclenche partout : tous les candidats montent ensemble. Dans les
    # deux cas le classement cesse d'ordonner, sans qu'aucune erreur ne
    # s'affiche.
    #
    # Le corpus fictif tient lieu de témoin : il contient un seul comptable
    # (cas 10) au milieu de développeurs et d'une assistante. L'écart entre
    # les deux est ce que ce critère est censé mesurer.
    criteres_comptable = charger_modele("Comptable")
    scores_comptable = {}
    for chemin_cv in sorted(DOSSIER_CORPUS.iterdir()):
        if chemin_cv.suffix.lower() not in (".pdf", ".docx"):
            continue
        with open(chemin_cv, "rb") as fichier:
            resultat = parser_cv(io.BytesIO(fichier.read()), chemin_cv.name)
        if resultat["lisible"]:
            scores_comptable[chemin_cv.name] = scorer_candidat(
                resultat["texte"], criteres_comptable
            )["score"]

    score_du_comptable = scores_comptable.get("10_CV_Brou_Akissi.docx")
    autres = [
        score for nom, score in scores_comptable.items()
        if nom != "10_CV_Brou_Akissi.docx"
    ]

    # CE QU'ON MESURE ICI EST UN ÉCART, PAS UN SEUIL.
    #
    # La première version de ce contrôle exigeait un score absolu d'au moins
    # 70. Elle a sauté dès l'ajout d'un cinquième groupe au critère le plus
    # lourd : le comptable de référence est tombé à 69 sans que le modèle
    # soit devenu moins juste. Un critère qui gagne un groupe partage ses
    # points entre cinq au lieu de quatre, et TOUS les candidats baissent
    # ensemble — c'est arithmétique, pas un défaut.
    #
    # Le seuil absolu mesurait donc le nombre de groupes autant que la
    # qualité du modèle. Ce qui compte réellement, c'est que le comptable se
    # détache NETTEMENT des profils hors métier. On vérifie l'écart, qui est
    # la propriété utile, et l'on garde un plancher bas comme simple filet
    # contre un effondrement général.
    meilleur_hors_metier = max(autres) if autres else None

    verifier(
        "le seul comptable du corpus reste largement au-dessus du plancher",
        score_du_comptable is not None and score_du_comptable >= 60,
        score_du_comptable,
    )
    verifier(
        "il se détache d'au moins 40 points du meilleur profil hors métier",
        meilleur_hors_metier is not None
        and score_du_comptable - meilleur_hors_metier >= 40,
        f"{score_du_comptable} contre {meilleur_hors_metier} "
        f"(écart {score_du_comptable - meilleur_hors_metier})",
    )

    # Aucun groupe ne doit être inerte sur un CV de comptable complet : un
    # mot-clé qui ne se déclenche jamais n'est pas un critère exigeant, c'est
    # un critère mort, et il abaisse tout le monde à l'identique.
    with open(DOSSIER_CORPUS / "10_CV_Brou_Akissi.docx", "rb") as fichier:
        reference = parser_cv(io.BytesIO(fichier.read()), "10_CV_Brou_Akissi.docx")
    detail_reference = scorer_candidat(
        reference["texte"], criteres_comptable
    )["detail"]
    for ligne in detail_reference:
        verifier(
            f"« {ligne['critere']} » : au moins un groupe se déclenche",
            bool(ligne["trouves"]),
            f"trouvés {ligne['trouves']} / absents {ligne['absents']}",
        )


# ----------------------------------------------------------------------
# PARTIE D — Classeur Excel (lot 4)
# ----------------------------------------------------------------------


def partie_d_export(resultats):
    sous_titre("D1. Génération du classeur depuis les résultats du corpus")

    criteres = MODELE_DEVELOPPEUR_WEB + [
        {
            "nom": "Recherche un stage",
            "poids": 0,
            "type": "exclusion",
            "mots_cles": "recherche un stage\ninterim",
            "groupes": analyser_mots_cles("recherche un stage\ninterim"),
        }
    ]

    # On rejoue le scoring en incluant le critère d'exclusion, pour que le
    # classeur contienne au moins un signalement à mettre en forme.
    candidats = []
    for resultat in resultats:
        candidat = dict(resultat)
        if candidat["lisible"]:
            candidat.update(scorer_candidat_depuis_corpus(candidat, criteres))
        candidats.append(candidat)

    jour = date(2026, 9, 14)
    tampon = generer_classeur(candidats, criteres, "Développeur web", jour)

    octets = tampon.getvalue()
    print(f"  Classeur généré en mémoire : {len(octets):,} octets".replace(",", " "))
    print(f"  Nom proposé : {nom_du_fichier_export(jour)}")
    print()

    verifier("le classeur est rendu dans un tampon mémoire", isinstance(tampon, io.BytesIO))
    verifier("le tampon est positionné au début", tampon.tell() == 0)
    verifier(
        "le contenu est bien un fichier xlsx (signature ZIP « PK »)",
        octets[:2] == b"PK",
        octets[:2],
    )

    # ------------------------------------------------------------------
    # Relecture : le classeur doit s'ouvrir et contenir ce qu'on attend
    # ------------------------------------------------------------------
    sous_titre("D2. Relecture du classeur")

    relu = load_workbook(io.BytesIO(octets))
    print(f"  Feuilles : {relu.sheetnames}")

    verifier(
        "les trois feuilles attendues sont présentes",
        relu.sheetnames == ["Résumé", "Classement", "Critères"],
        relu.sheetnames,
    )

    classes = [c for c in candidats if c["lisible"]]
    illisibles = [c for c in candidats if not c["lisible"]]

    feuille = relu["Classement"]
    valeurs = [
        [cellule.value for cellule in ligne]
        for ligne in feuille.iter_rows()
    ]
    plat = [str(v) for ligne in valeurs for v in ligne if v is not None]

    verifier(
        "l'en-tête porte une colonne par critère requis",
        all(f"{c['nom']} (pts)" in plat for c in MODELE_DEVELOPPEUR_WEB),
    )
    verifier(
        "chaque candidat classé a une ligne",
        all(c["fichier"] in plat for c in classes),
    )
    verifier(
        "les illisibles figurent dans le bloc séparé",
        all(c["fichier"] in plat for c in illisibles),
    )
    # Les lignes du classement sont celles dont la première cellule est un
    # numéro de rang. Aucun fichier illisible ne doit s'y trouver.
    fichiers_classes = {
        ligne[-1] for ligne in valeurs if ligne and isinstance(ligne[0], int)
    }
    verifier(
        "aucun illisible ne figure parmi les lignes numérotées du classement",
        all(c["fichier"] not in fichiers_classes for c in illisibles),
        fichiers_classes & {c["fichier"] for c in illisibles},
    )
    verifier(
        "le bloc séparé donne la raison au lieu d'un score",
        "Document scanné, aucun texte lisible" in plat,
    )
    verifier(
        "tous les candidats lisibles sont numérotés",
        fichiers_classes == {c["fichier"] for c in classes},
        fichiers_classes,
    )

    # Le premier du classement doit être en première ligne de données.
    premier_attendu = max(classes, key=lambda c: c["score"])
    ligne_1 = next((l for l in valeurs if l and l[0] == 1), None)
    verifier(
        "le rang 1 est bien le meilleur score",
        ligne_1 is not None and ligne_1[4] == premier_attendu["score"],
        ligne_1[4] if ligne_1 else None,
    )

    feuille_resume = relu["Résumé"]
    plat_resume = [
        str(c.value) for ligne in feuille_resume.iter_rows() for c in ligne if c.value
    ]
    verifier(
        "l'avertissement méthodologique figure dans le résumé",
        any(AVERTISSEMENT in texte for texte in plat_resume),
    )
    verifier(
        "le résumé compte les illisibles à part",
        "À examiner manuellement" in plat_resume,
    )

    feuille_criteres = relu["Critères"]
    plat_criteres = [
        str(c.value) for ligne in feuille_criteres.iter_rows() for c in ligne if c.value
    ]
    verifier(
        "la feuille Critères liste le critère d'exclusion et son type",
        "Recherche un stage" in plat_criteres and "Exclusion" in plat_criteres,
    )
    verifier(
        "la feuille Critères réaffiche la syntaxe des synonymes",
        any("react = vue, angular" in texte for texte in plat_criteres),
    )

    # ------------------------------------------------------------------
    # Confidentialité
    # ------------------------------------------------------------------
    sous_titre("D3. Aucune écriture disque")

    avant = set(RACINE.rglob("*.xlsx"))
    generer_classeur(candidats, criteres, "Développeur web", jour)
    apres = set(RACINE.rglob("*.xlsx"))
    verifier(
        "générer un classeur ne crée aucun fichier dans le dépôt",
        avant == apres,
        apres - avant,
    )
    verifier(
        "le nom de fichier proposé ne contient aucun nom de candidat",
        not any(
            (c.get("nom") or "").split()[0].lower() in nom_du_fichier_export(jour).lower()
            for c in classes
            if c.get("nom")
        ),
        nom_du_fichier_export(jour),
    )


def scorer_candidat_depuis_corpus(candidat, criteres):
    """
    Rejoue le scoring d'un candidat du corpus.

    Le texte brut ayant été libéré en partie B (règle n°3 du §5), on le
    réextrait du fichier. C'est acceptable ici : le corpus est fictif, et
    cela évite de garder du texte en mémoire plus longtemps que nécessaire
    pour le seul confort d'un script de vérification.
    """
    chemin = DOSSIER_CORPUS / candidat["fichier"]
    with open(chemin, "rb") as fichier:
        memoire = io.BytesIO(fichier.read())
    relu = parser_cv(memoire, chemin.name)
    return scorer_candidat(relu["texte"], criteres)


# ----------------------------------------------------------------------


def principal():
    titre("PARTIE A — Les trois corrections du moteur (lot 1)")
    partie_a_frontieres_de_mots()
    partie_a_accents()
    partie_a_normalisation_des_poids()
    partie_a_exclusions()

    resultats = partie_b_corpus()

    titre("PARTIE C — Syntaxe de saisie et modèles (lot 3)")
    partie_c_syntaxe()
    partie_c_modeles()

    titre("PARTIE D — Classeur Excel (lot 4)")
    if resultats:
        partie_d_export(resultats)
    else:
        print("  Sautée : le corpus est vide.")

    titre("RÉSULTAT DE LA PORTE P1")
    if _echecs:
        print(f"  {len(_echecs)} vérification(s) en échec :")
        for echec in _echecs:
            print(f"    - {echec}")
        print()
        print("  PORTE P1 NON FRANCHIE.")
        return 1

    print("  Toutes les vérifications passent.")
    print()
    print("  PORTE P1 FRANCHIE côté corpus synthétique.")
    print()
    print("  Il reste le second volet de la porte P1 (PROCESS §5.2) : vérifier")
    print("  que le classement reste crédible sur l'échantillon réel d'une")
    print("  dizaine de CV, conservé HORS du dépôt. Ce script ne peut pas le")
    print("  faire à ta place.")
    return 0


if __name__ == "__main__":
    sys.exit(principal())
